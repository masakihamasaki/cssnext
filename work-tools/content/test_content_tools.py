import json
import os
import tempfile
import unittest

from pptx import Presentation

import manuscript_lint
import narration
from readings import Converter, number_to_kanji


class ReadingsTest(unittest.TestCase):
    def test_numbers(self):
        cases = {0: "ゼロ", 10: "十", 15: "十五", 100: "百", 1873: "千八百七十三",
                 2027: "二千二十七", 10000: "一万", 15000000: "千五百万",
                 120000000: "一億二千万", 100010: "十万十"}
        for n, want in cases.items():
            self.assertEqual(number_to_kanji(n), want, n)

    def test_text(self):
        c = Converter()
        self.assertEqual(c.convert("2,000人から3,000人"), "二千人から三千人")
        self.assertEqual(c.convert("AIとM&A"), "エーアイとエムアンドエー")
        self.assertEqual(c.convert("15.5%増"), "十五点五パーセント増")
        self.assertEqual(c.convert("10〜20件"), "十から二十件")
        self.assertEqual(c.convert("ＤＸ推進"), "ディーエックス推進")
        self.assertEqual(c.convert("XYZ社"), "エックスワイゼット社")

    def test_unknown_words_are_kept_and_reported(self):
        c = Converter()
        self.assertEqual(c.convert("Higgsfieldを使う"), "Higgsfieldを使う")
        self.assertIn("Higgsfield", c.unknown)

    def test_private_dictionary_overrides(self):
        c = Converter({"foo": "フー"})
        self.assertEqual(c.convert("Foo"), "フー")


class NarrationTest(unittest.TestCase):
    def test_split(self):
        text = "# 見出し\n\n一文目です。二文目ですか？\n- 箇条書き\n---\n[リンク](http://x)の文。"
        self.assertEqual(narration.split_sentences(text),
                         ["見出し", "一文目です。", "二文目ですか？", "箇条書き", "リンクの文。"])

    def test_split_long(self):
        s = "あ" * 30 + "、" + "い" * 30 + "。"
        self.assertEqual(len(narration.split_long(s, 40)), 2)
        self.assertEqual(narration.split_long("う" * 80, 40), ["う" * 80])

    def test_variant(self):
        self.assertEqual(narration.apply_variant(list("abcd"), 2, "X"), list("abXcdX"))

    def test_end_to_end(self):
        with tempfile.TemporaryDirectory() as d:
            src = os.path.join(d, "s.md")
            with open(src, "w", encoding="utf-8") as f:
                f.write("2027年に始めます。AIを使います。三つ目です。\n")
            open(os.path.join(d, "ABC_004_v1_20260101.pptx"), "w").close()
            narration.main([src, "--series", "ABC", "--out-dir", d, "--date", "20261010",
                            "--per-file", "2", "--variant", "ABV:2:定型文です。"])
            names = sorted(f for f in os.listdir(d) if f.endswith(".pptx"))
            self.assertIn("ABC_005_v1_20261010.pptx", names)
            self.assertIn("ABC_006_v1_20261010.pptx", names)
            self.assertIn("ABV_001_v1_20261010.pptx", names)
            prs = Presentation(os.path.join(d, "ABC_005_v1_20261010.pptx"))
            texts = [s.shapes[0].text_frame.text for s in prs.slides]
            self.assertEqual(texts, ["二千二十七年に始めます。", "エーアイを使います。"])
            self.assertEqual(prs.slides[0].notes_slide.notes_text_frame.text, texts[0])
            with open(os.path.join(d, "ABC_report_20261010.json"), encoding="utf-8") as f:
                self.assertEqual(json.load(f)["unconverted_words"], [])


class LintTest(unittest.TestCase):
    def test_check(self):
        spec = {"min_chars": 50, "required": ["出典"], "banned": ["必ず"], "numbers_need_source": True}
        n, issues = manuscript_lint.check("売上は100万円です。必ず伸びます。", spec)
        joined = "\n".join(issues)
        self.assertIn("字数不足", joined)
        self.assertIn("必要な要素なし: 出典", joined)
        self.assertIn("禁止表現「必ず」", joined)
        self.assertIn("数字の出典なし", joined)
        _, ok = manuscript_lint.check("統計によると100万円です。", {"numbers_need_source": True})
        self.assertEqual(ok, [])


if __name__ == "__main__":
    unittest.main()
