#!/usr/bin/env python3
"""Turn a manuscript into one-sentence narration slides (.pptx) for Google Vids.

    python3 narration.py source.md --series ABC --out-dir out/
    python3 narration.py source.md --series ABC --variant ABV:3:定型文です。

Each slide holds one sentence rewritten for text-to-speech (see readings.py);
the same text goes into the speaker notes. Files are named
<SERIES>_<NNN>_v1_<YYYYMMDD>.pptx and numbered after the highest number
already in --out-dir. A JSON report lists replacements, words that could not
be converted, and sentences that are too long to read comfortably.
"""

import argparse
import datetime
import html
import json
import os
import re
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from readings import Converter, load_dictionaries  # noqa: E402

SLIDE_W, SLIDE_H = Emu(12192000), Emu(6858000)  # 16:9
END_PUNCT = "。！？!?"


def read_source(path):
    if path.endswith(".docx"):
        import docx
        return "\n".join(p.text for p in docx.Document(path).paragraphs)
    with open(path, encoding="utf-8") as f:
        text = f.read()
    if path.endswith((".html", ".htm")):
        text = re.sub(r"(?is)<(script|style)\b.*?</\1>", "", text)
        text = re.sub(r"(?i)<br\s*/?>|</(p|h\d|li|div)>", "\n", text)
        text = html.unescape(re.sub(r"<[^>]+>", "", text))
    return text


def clean_line(line):
    line = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", line)          # images
    line = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", line)      # links
    line = re.sub(r"^\s{0,3}(#{1,6}|>|[-*+]|\d+[.)])\s+", "", line)
    line = re.sub(r"[*_`]{1,3}", "", line)
    return line.strip()


def split_sentences(text):
    """One entry per sentence; lines without end punctuation are headings."""
    out = []
    for raw in text.splitlines():
        line = clean_line(raw)
        if not line or re.fullmatch(r"[-=_*]{3,}", line):
            continue
        parts = re.findall(r"[^%s]+[%s」』）)]*" % (END_PUNCT, END_PUNCT), line)
        out.extend(p.strip() for p in parts if p.strip())
    return out


def split_long(sentence, max_chars):
    """Split at the 、 closest to the middle; give up if there is none."""
    if len(sentence) <= max_chars:
        return [sentence]
    commas = [i for i, c in enumerate(sentence) if c == "、"]
    if not commas:
        return [sentence]
    mid = len(sentence) // 2
    cut = min(commas, key=lambda i: abs(i - mid)) + 1
    return split_long(sentence[:cut], max_chars) + split_long(sentence[cut:], max_chars)


def apply_variant(slides, every, text):
    out = []
    for i, s in enumerate(slides, 1):
        out.append(s)
        if i % every == 0:
            out.append(text)
    return out


def next_number(out_dir, series):
    pat = re.compile(r"^%s_(\d+)" % re.escape(series))
    nums = [int(m.group(1)) for f in os.listdir(out_dir) if (m := pat.match(f))]
    return max(nums, default=0) + 1


def build_pptx(sentences, path, font):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    layout = prs.slide_layouts[6]  # blank
    margin = Emu(int(SLIDE_W * 0.08))
    for s in sentences:
        slide = prs.slides.add_slide(layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        box = slide.shapes.add_textbox(margin, 0, SLIDE_W - 2 * margin, SLIDE_H)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = s
        run.font.size = Pt(44 if len(s) <= 30 else 36 if len(s) <= 50 else 30)
        run.font.name = font
        run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        slide.notes_slide.notes_text_frame.text = s
    prs.save(path)


def chunk(items, size):
    return [items[i:i + size] for i in range(0, len(items), size)] or [[]]


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("source", help=".md / .txt / .html / .docx")
    ap.add_argument("--series", required=True, help="file name prefix, e.g. ABC")
    ap.add_argument("--out-dir", default=".")
    ap.add_argument("--start", type=int, help="first number (default: next free number in out-dir)")
    ap.add_argument("--per-file", type=int, default=30, help="slides per file (default 30)")
    ap.add_argument("--max-chars", type=int, default=60, help="split sentences longer than this at 、")
    ap.add_argument("--variant", action="append", default=[], metavar="SERIES:EVERY:TEXT",
                    help="also build SERIES with TEXT inserted after every EVERY slides")
    ap.add_argument("--dict", action="append", default=[], help="extra reading dictionary JSON (private terms)")
    ap.add_argument("--font", default="Noto Sans JP")
    ap.add_argument("--date", default=datetime.date.today().strftime("%Y%m%d"))
    args = ap.parse_args(argv)

    os.makedirs(args.out_dir, exist_ok=True)
    conv = Converter(*load_dictionaries(args.dict))
    sentences, long_lines = [], []
    for s in split_sentences(read_source(args.source)):
        spoken = conv.convert(s)
        pieces = split_long(spoken, args.max_chars)
        long_lines += [p for p in pieces if len(p) > args.max_chars]
        sentences += pieces

    plans = [(args.series, sentences)]
    for spec in args.variant:
        name, every, text = spec.split(":", 2)
        plans.append((name, apply_variant(sentences, int(every), conv.convert(text))))

    files = []
    for series, slides in plans:
        number = args.start if args.start else next_number(args.out_dir, series)
        for part in chunk(slides, args.per_file):
            name = "%s_%03d_v1_%s.pptx" % (series, number, args.date)
            build_pptx(part, os.path.join(args.out_dir, name), args.font)
            files.append({"file": name, "slides": len(part)})
            number += 1

    report = {
        "files": files,
        "replacements": conv.replaced,
        "unconverted_words": sorted(conv.unknown),
        "too_long": long_lines,
    }
    with open(os.path.join(args.out_dir, "%s_report_%s.json" % (args.series, args.date)), "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
